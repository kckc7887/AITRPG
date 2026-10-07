import hashlib
import io
import locale
import os
import re
import shutil
import stat
import subprocess
import threading
import zipfile
from dataclasses import dataclass
from dataclasses import field
from pathlib import Path
from pathlib import PurePosixPath

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfReader

from aitrpg.domain.models import ReviewIssue
from aitrpg.domain.models import ScenarioAsset
from aitrpg.domain.models import SourceBlock

MAX_FILE_BYTES = 128 * 1024 * 1024
MAX_TOTAL_BYTES = 512 * 1024 * 1024
MAX_FILES = 5000
MAX_ARCHIVE_DEPTH = 4
MAX_TEXT_CHARS = 800000
MAX_IMAGE_PIXELS = 40000000
IMAGE_SUFFIXES = {'.png', '.jpg', '.jpeg', '.webp', '.gif', '.bmp'}
DEVICE_NAMES = {
    'con',
    'prn',
    'aux',
    'nul',
    'conin$',
    'conout$',
    *(f'com{number}' for number in '123456789¹²³'),
    *(f'lpt{number}' for number in '123456789¹²³'),
}


class AttachmentError(ValueError):
    pass


@dataclass
class ExtractionBudget:
    total_bytes: int = 0
    files: int = 0

    def consume(self, size: int) -> None:
        if size < 0 or size > MAX_FILE_BYTES:
            raise AttachmentError('附件成员超过单文件大小限制')
        self.total_bytes += size
        self.files += 1
        if self.total_bytes > MAX_TOTAL_BYTES or self.files > MAX_FILES:
            raise AttachmentError('附件解包总量超过限制')


@dataclass
class ParsedAttachment:
    blocks: list[SourceBlock] = field(default_factory=list)
    assets: list[ScenarioAsset] = field(default_factory=list)
    issues: list[ReviewIssue] = field(default_factory=list)
    source_dir: Path | None = None


def safe_member_path(directory: Path, name: str) -> Path:
    normalized = name.replace('\\', '/')
    relative = PurePosixPath(normalized)
    if (
        relative.is_absolute()
        or not relative.parts
        or any(part in {'.', '..'} or ':' in part for part in relative.parts)
        or any(
            part.endswith(('.', ' '))
            or part.split('.', 1)[0].rstrip(' ').lower() in DEVICE_NAMES
            or any(char in '<>"|?*' for char in part)
            for part in relative.parts
        )
        or any(ord(char) < 32 for char in normalized)
    ):
        raise AttachmentError('附件包含不安全的路径')
    target = directory.joinpath(*relative.parts).resolve()
    if not target.is_relative_to(directory.resolve()):
        raise AttachmentError('附件路径越出目标目录')
    return target


def _zip_name(item: zipfile.ZipInfo) -> str:
    if item.flag_bits & 0x800:
        return item.filename
    try:
        return item.filename.encode('cp437').decode('gb18030')
    except (UnicodeEncodeError, UnicodeDecodeError):
        return item.filename


def _decode_output(value: bytes) -> str:
    for encoding in ('utf-8', locale.getpreferredencoding(False), 'gb18030'):
        try:
            return value.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise AttachmentError('无法确定压缩文件名的编码')


def _bsdtar() -> str:
    configured = os.environ.get('AITRPG_BSDTAR', '')
    candidates = [configured, shutil.which('bsdtar'), shutil.which('tar')]
    for candidate in candidates:
        if not candidate:
            continue
        result = subprocess.run(
            [candidate, '--version'],
            capture_output=True,
            timeout=10,
            check=False,
        )
        if result.returncode == 0 and b'bsdtar' in result.stdout.lower():
            return candidate
    raise AttachmentError(
        '读取 RAR 需要 bsdtar；设置 AITRPG_BSDTAR 或先自行解包'
    )


def _rar_members(path: Path, executable: str) -> list[tuple[str, int]]:
    result = subprocess.run(
        [executable, '-tvf', str(path)],
        capture_output=True,
        timeout=30,
        check=False,
    )
    if result.returncode:
        raise AttachmentError('RAR 列表读取失败，可能加密或格式不支持')
    members = []
    for line in _decode_output(result.stdout).splitlines():
        if not line.strip():
            continue
        fields = line.split(maxsplit=8)
        if len(fields) != 9 or fields[0][0] not in {'-', 'd'}:
            raise AttachmentError('RAR 包含链接或无法安全识别的成员')
        if fields[0][0] == 'd':
            continue
        try:
            size = int(fields[4])
        except ValueError as error:
            raise AttachmentError('RAR 成员大小不可识别') from error
        members.append((fields[8], size))
    return members


def _rar_bytes(path: Path, name: str, executable: str) -> bytes:
    process = subprocess.Popen(
        [executable, '-xOf', str(path), '--', name],
        stdout=subprocess.PIPE,
        stderr=subprocess.DEVNULL,
    )
    timer = threading.Timer(30, process.kill)
    timer.start()
    try:
        if process.stdout is None:
            raise AttachmentError('RAR 解压输出不可用')
        payload = process.stdout.read(MAX_FILE_BYTES + 1)
        if len(payload) > MAX_FILE_BYTES:
            raise AttachmentError('RAR 实际解压大小超过限制')
        if process.wait(timeout=5):
            raise AttachmentError('RAR 成员解压失败')
        return payload
    finally:
        timer.cancel()
        if process.poll() is None:
            process.kill()
            process.wait(timeout=5)
        if process.stdout is not None:
            process.stdout.close()


def unpack_archive(
    path: Path,
    directory: Path,
    budget: ExtractionBudget | None = None,
    depth: int = 0,
) -> list[Path]:
    if depth > MAX_ARCHIVE_DEPTH:
        raise AttachmentError('嵌套压缩文件层数超过限制')
    budget = budget or ExtractionBudget()
    directory.mkdir(parents=True, exist_ok=True)
    paths = []
    names = set()
    if path.suffix.lower() == '.zip':
        with zipfile.ZipFile(path) as archive:
            for item in archive.infolist():
                name = _zip_name(item)
                target = safe_member_path(directory, name)
                mode = item.external_attr >> 16
                if stat.S_ISLNK(mode):
                    raise AttachmentError('ZIP 不允许符号链接')
                if item.is_dir():
                    continue
                if item.flag_bits & 1:
                    raise AttachmentError('不支持加密 ZIP')
                identity = str(target).casefold()
                if identity in names:
                    raise AttachmentError('ZIP 包含冲突成员路径')
                names.add(identity)
                budget.consume(item.file_size)
                with archive.open(item) as stream:
                    payload = stream.read(MAX_FILE_BYTES + 1)
                if len(payload) != item.file_size:
                    raise AttachmentError('ZIP 成员实际大小与声明不符')
                target.parent.mkdir(parents=True, exist_ok=True)
                target.write_bytes(payload)
                paths.append(target)
    elif path.suffix.lower() == '.rar':
        executable = _bsdtar()
        for name, size in _rar_members(path, executable):
            target = safe_member_path(directory, name)
            identity = str(target).casefold()
            if identity in names:
                raise AttachmentError('RAR 包含冲突成员路径')
            names.add(identity)
            budget.consume(size)
            payload = _rar_bytes(path, name, executable)
            if len(payload) != size:
                raise AttachmentError('RAR 成员实际大小与声明不符')
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_bytes(payload)
            paths.append(target)
    else:
        raise AttachmentError('不支持该压缩格式')
    result = []
    for member in paths:
        if member.suffix.lower() in {'.zip', '.rar'}:
            result.extend(
                unpack_archive(
                    member,
                    member.parent / (member.name + '.contents'),
                    budget,
                    depth + 1,
                )
            )
        else:
            result.append(member)
    return result


def _add_block(
    parsed: ParsedAttachment,
    file: str,
    locator: str,
    text: str,
    kind: str = 'text',
) -> None:
    text = text.strip()
    offset = 0
    while offset < len(text):
        end = min(len(text), offset + 6000)
        if end < len(text):
            boundary = text.rfind('\n', offset + 2000, end)
            if boundary != -1:
                end = boundary + 1
        part_locator = locator
        if len(text) > 6000:
            part_locator += f':chars:{offset}-{end}'
        parsed.blocks.append(
            SourceBlock(
                file=file,
                locator=part_locator,
                text=text[offset:end],
                kind=kind,
            )
        )
        offset = end


def _add_image(
    parsed: ParsedAttachment,
    payload: bytes,
    name: str,
    file: str,
    locator: str,
    directory: Path,
) -> None:
    with Image.open(io.BytesIO(payload)) as image:
        if image.width * image.height > MAX_IMAGE_PIXELS:
            raise AttachmentError('图片解码尺寸超过限制')
        image.load()
        digest = hashlib.sha256(payload).hexdigest()
        asset_path = directory / 'assets' / (digest + '.png')
        asset_path.parent.mkdir(parents=True, exist_ok=True)
        image.convert('RGBA').save(asset_path)
    block = SourceBlock(
        file=file,
        locator=locator,
        kind='image',
        text=name,
    )
    parsed.blocks.append(block)
    if any(asset.path == f'assets/{digest}.png' for asset in parsed.assets):
        existing = next(
            asset
            for asset in parsed.assets
            if asset.path == f'assets/{digest}.png'
        )
        existing.source_ids.append(block.id)
        return
    parsed.assets.append(
        ScenarioAsset(
            name=name,
            path=f'assets/{digest}.png',
            mime_type='image/png',
            is_map=bool(re.search(r'地图|平面图|map', name, re.IGNORECASE)),
            source_ids=[block.id],
        )
    )


def _docx_blocks(
    path: Path,
    file: str,
    parsed: ParsedAttachment,
    directory: Path,
) -> None:
    _validate_office_archive(path)
    document = Document(path)
    paragraph_index = 0
    table_index = 0
    for element in document.element.body:
        if element.tag.endswith('}p'):
            paragraph_index += 1
            paragraph = Paragraph(element, document)
            _add_block(
                parsed, file, f'paragraph:{paragraph_index}', paragraph.text
            )
        elif element.tag.endswith('}tbl'):
            table_index += 1
            table = Table(element, document)
            seen = set()
            for row_index, row in enumerate(table.rows, start=1):
                for cell_index, cell in enumerate(row.cells, start=1):
                    if cell._tc in seen:
                        continue
                    seen.add(cell._tc)
                    locator = (
                        f'table:{table_index}:row:{row_index}'
                        f':cell:{cell_index}'
                    )
                    _add_block(parsed, file, locator, cell.text, 'table_cell')
    with zipfile.ZipFile(path) as archive:
        for item in archive.infolist():
            if item.is_dir() or not item.filename.startswith('word/media/'):
                continue
            if Path(item.filename).suffix.lower() not in IMAGE_SUFFIXES:
                parsed.issues.append(
                    ReviewIssue(
                        message=f'{file} 包含未支持的图片：{item.filename}',
                    )
                )
                continue
            if item.file_size > MAX_FILE_BYTES:
                raise AttachmentError('DOCX 图片超过大小限制')
            _add_image(
                parsed,
                archive.read(item),
                item.filename,
                file,
                item.filename,
                directory,
            )


def _pdf_blocks(
    path: Path,
    file: str,
    parsed: ParsedAttachment,
    directory: Path,
) -> None:
    reader = PdfReader(path)
    if reader.is_encrypted:
        raise AttachmentError('暂不支持加密 PDF')
    for index, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ''
        _add_block(parsed, file, f'page:{index}', text)
        if not text.strip():
            parsed.issues.append(
                ReviewIssue(
                    message=f'{file} 第 {index} 页缺少文本层，请复核图片页',
                )
            )
        for image_index, image in enumerate(page.images, start=1):
            locator = f'page:{index}:image:{image_index}'
            try:
                _add_image(
                    parsed, image.data, image.name, file, locator, directory
                )
            except (OSError, ValueError) as error:
                parsed.issues.append(
                    ReviewIssue(
                        message=f'{file} {locator} 图片无法解码：{error}',
                    )
                )


def _xlsx_blocks(path: Path, file: str, parsed: ParsedAttachment) -> None:
    _validate_office_archive(path)
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        sheets = [
            sheet
            for sheet in workbook
            if sheet.title in {'人物卡', '角色卡', '调查员'}
        ]
        if not sheets:
            sheets = [workbook.worksheets[0]]
        omitted = [sheet.title for sheet in workbook if sheet not in sheets]
        if omitted:
            parsed.issues.append(
                ReviewIssue(
                    message=(
                        f'{file} 仅索引角色主表，辅助表和隐藏规则表'
                        '完整保留在原始工作簿中，未执行公式'
                    ),
                )
            )
        for sheet in sheets:
            if sheet.max_row > 10000 or sheet.max_column > 256:
                raise AttachmentError('工作表范围过大，请先清理异常格式')
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    locator = f'sheet:{sheet.title}:cell:{cell.coordinate}'
                    kind = (
                        'formula'
                        if cell.data_type == 'f'
                        else 'character_cell'
                    )
                    _add_block(parsed, file, locator, str(cell.value), kind)
    finally:
        workbook.close()


def _normalized_text(blocks: list[SourceBlock]) -> set[str]:
    text = ''.join(
        re.findall(
            r'[\w\u4e00-\u9fff]', ''.join(block.text for block in blocks)
        )
    )
    return {text[index : index + 5] for index in range(len(text) - 4)}


def _validate_office_archive(path: Path) -> None:
    budget = ExtractionBudget()
    with zipfile.ZipFile(path) as archive:
        for item in archive.infolist():
            if not item.is_dir():
                budget.consume(item.file_size)


def parse_attachment(path: Path, directory: Path) -> ParsedAttachment:
    path = path.resolve()
    if not path.exists() or path.is_symlink():
        raise AttachmentError('来源文件不存在或为符号链接')
    raw_dir = directory / 'raw'
    raw_dir.mkdir(parents=True, exist_ok=True)
    parsed = ParsedAttachment(source_dir=raw_dir)
    budget = ExtractionBudget()
    if path.is_dir():
        files = []
        for source in sorted(path.rglob('*')):
            if source.is_symlink():
                raise AttachmentError('导入目录不允许符号链接')
            if not source.is_file():
                continue
            budget.consume(source.stat().st_size)
            target = safe_member_path(
                raw_dir,
                source.relative_to(path).as_posix(),
            )
            target.parent.mkdir(parents=True, exist_ok=True)
            shutil.copyfile(source, target)
            if target.suffix.lower() in {'.zip', '.rar'}:
                files.extend(
                    unpack_archive(
                        target,
                        target.parent / (target.name + '.contents'),
                        budget,
                    )
                )
            else:
                files.append(target)
    else:
        budget.consume(path.stat().st_size)
        copied = raw_dir / path.name
        shutil.copyfile(path, copied)
        if path.suffix.lower() in {'.zip', '.rar'}:
            files = unpack_archive(copied, raw_dir / 'extracted', budget)
        else:
            files = [copied]
    by_file = {}
    for source in files:
        file = source.relative_to(directory).as_posix()
        partial = ParsedAttachment()
        suffix = source.suffix.lower()
        try:
            if suffix == '.docx':
                _docx_blocks(source, file, partial, directory)
            elif suffix == '.pdf':
                _pdf_blocks(source, file, partial, directory)
            elif suffix == '.xlsx':
                _xlsx_blocks(source, file, partial)
            elif suffix in IMAGE_SUFFIXES:
                _add_image(
                    partial,
                    source.read_bytes(),
                    source.name,
                    file,
                    'image',
                    directory,
                )
            elif suffix in {'.txt', '.md'}:
                _add_block(
                    partial,
                    file,
                    'text',
                    _decode_output(source.read_bytes()),
                    'metadata' if source.stem.lower() == 'readme' else 'text',
                )
            elif suffix not in {'.json'}:
                partial.issues.append(
                    ReviewIssue(
                        message=f'保留但未解析的文件：{file}',
                    )
                )
        except AttachmentError:
            raise
        except Exception as error:
            partial.issues.append(
                ReviewIssue(
                    message=f'无法读取 {file}：{type(error).__name__}',
                    severity='blocker',
                )
            )
        by_file[source] = partial
    for source, partial in by_file.items():
        if source.suffix.lower() == '.pdf':
            matching = source.with_suffix('.docx')
            if matching in by_file:
                first = _normalized_text(partial.blocks)
                second = _normalized_text(by_file[matching].blocks)
                denominator = max(1, max(len(first), len(second)))
                overlap = len(first & second) / denominator
                if overlap >= 0.8:
                    parsed.issues.append(
                        ReviewIssue(
                            message=(
                                f'{source.name} 与 DOCX 重复，采用 DOCX 结构'
                            ),
                        )
                    )
                    continue
                parsed.issues.append(
                    ReviewIssue(
                        message=(
                            f'{source.name} 与同名 DOCX 内容不同，需确认版本'
                        ),
                        severity='blocker',
                    )
                )
        parsed.blocks.extend(partial.blocks)
        parsed.assets.extend(partial.assets)
        parsed.issues.extend(partial.issues)
    if sum(len(block.text) for block in parsed.blocks) > MAX_TEXT_CHARS:
        raise AttachmentError('来源文字总量超过单次导入限制')
    return parsed
