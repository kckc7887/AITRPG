import xml.etree.ElementTree as ElementTree
from pathlib import Path
from typing import Any
from zipfile import BadZipFile
from zipfile import ZipFile

from openpyxl import load_workbook

from aitrpg.domain.models import Character
from aitrpg.domain.rules import BACKGROUND_FIELDS
from aitrpg.domain.rules import SKILL_BASES
from aitrpg.domain.rules import base_skill
from aitrpg.domain.rules import derived_character
from aitrpg.domain.rules import normalize_skill

ATTRIBUTE_CELLS = {
    'STR': 'U3',
    'CON': 'U5',
    'SIZ': 'U7',
    'DEX': 'AA3',
    'APP': 'AA5',
    'INT': 'AA7',
    'POW': 'AG3',
    'EDU': 'AG5',
}
BACKGROUND_CELLS = {
    'personal_description': 'AA61',
    'ideology_beliefs': 'AA63',
    'significant_people': 'AA65',
    'meaningful_locations': 'AA67',
    'treasured_possessions': 'AA69',
    'traits': 'AA71',
    'injuries_scars': 'AA73',
    'phobias_manias': 'AA75',
    'life_story': 'W77',
}
CONDITION_NAMES = {
    '重伤': 'major_wound',
    '昏迷': 'unconscious',
    '濒死': 'dying',
    '死亡': 'dead',
    '临时疯狂': 'temporary_insanity',
    '不定性疯狂': 'indefinite_insanity',
    '永久疯狂': 'permanent_insanity',
}


def _safe_archive(path: Path) -> None:
    if path.suffix.lower() != '.xlsx' or path.stat().st_size > 20 * 1024**2:
        raise ValueError('仅支持不超过20MiB的xlsx角色卡')
    try:
        with ZipFile(path) as archive:
            entries = archive.infolist()
            if (
                len(entries) > 2000
                or sum(entry.file_size for entry in entries) > 100 * 1024**2
            ):
                raise ValueError('角色卡解压后过大')
            if any('vba' in entry.filename.lower() for entry in entries):
                raise ValueError('不接受包含宏的角色卡')
    except BadZipFile as exc:
        raise ValueError('角色卡不是有效的xlsx文件') from exc


def import_character_xlsx(path: str | Path, actor_id: str) -> Character:
    source_path = Path(path)
    _safe_archive(source_path)
    formulas = load_workbook(source_path, data_only=False, keep_links=False)
    values = load_workbook(source_path, data_only=True, keep_links=False)
    try:
        if '人物卡' not in formulas.sheetnames:
            raise ValueError('未识别CY23.5人物卡工作表')
        sheet = formulas['人物卡']
        cached = values['人物卡']
        warnings = []

        def cell(address: str) -> Any:
            original = sheet[address]
            if original.data_type == 'f':
                result = cached[address].value
                if result is None:
                    warnings.append(f'{address}缺少公式缓存，需人工核对')
                return result
            return original.value

        def text(address: str, default: str = '') -> str:
            value = cell(address)
            return str(value).strip() if value is not None else default

        def number(address: str, default: int | None = None) -> int:
            value = cell(address)
            if value in (None, '', '——', '—'):
                if default is None:
                    raise ValueError(f'角色卡{address}缺少数值')
                return default
            if isinstance(value, bool):
                raise ValueError(f'角色卡{address}不是有效整数')
            try:
                result = float(value)
            except (TypeError, ValueError) as exc:
                raise ValueError(f'角色卡{address}不是有效整数') from exc
            if not result.is_integer():
                raise ValueError(f'角色卡{address}须填写整数')
            return int(result)

        attributes = {
            name: number(address) for name, address in ATTRIBUTE_CELLS.items()
        }
        if any(value < 1 for value in attributes.values()):
            raise ValueError('空白模板不能作为已完成角色导入')
        skills = dict(SKILL_BASES)
        skills['母语'] = attributes['EDU']
        skills['闪避'] = attributes['DEX'] // 2
        marks = []
        imported_allocations = {}
        for name_col, subtype_col, columns, mark_col in (
            ('F', 'H', ('J', 'L', 'N', 'P'), 'B'),
            ('AB', 'AD', ('AF', 'AH', 'AJ', 'AL'), 'X'),
        ):
            for row in range(16, 50):
                name = text(f'{name_col}{row}').replace(' Ω', '')
                if not name:
                    continue
                subtype = text(f'{subtype_col}{row}')
                family = name.rstrip('：:①②③')
                if family in (
                    '技艺',
                    '科学',
                    '格斗',
                    '射击',
                    '外语',
                    '驾驶',
                    '生存',
                    '学识',
                    '自定义技能',
                ):
                    if not subtype:
                        continue
                    name = family + '：' + subtype
                name = normalize_skill(name)
                try:
                    initial = base_skill(name, attributes)
                except ValueError:
                    initial = number(f'{columns[0]}{row}', 0)
                    warnings.append(f'保留自定义技能{name}，由主持裁定')
                if sheet[f'{columns[0]}{row}'].data_type != 'f':
                    initial = number(f'{columns[0]}{row}', initial)
                increments = [number(f'{col}{row}', 0) for col in columns[1:]]
                total = initial + sum(increments)
                if total < 0:
                    raise ValueError(f'技能{name}总值不能为负数')
                skills[name] = total
                imported_allocations[name] = {
                    'base': initial,
                    'growth': increments[0],
                    'occupation': increments[1],
                    'interest': increments[2],
                }
                if text(f'{mark_col}{row}') == '☑':
                    marks.append(name)
        background = {name: '无' for name in BACKGROUND_FIELDS}
        for name, address in BACKGROUND_CELLS.items():
            background[name] = text(address, '无')
        for row, name in (
            (61, 'personal_description'),
            (63, 'ideology_beliefs'),
            (65, 'significant_people'),
            (67, 'meaningful_locations'),
            (69, 'treasured_possessions'),
            (71, 'traits'),
        ):
            if text(f'AR{row}') == '☑':
                background['key_connection'] = name
                break
        background['tomes_spells_artifacts'] = (
            '\n'.join(
                text(f'Y{row}') for row in range(114, 118) if text(f'Y{row}')
            )
            or '无'
        )
        experiences = [
            f'{text(f"B{row}")}：{text(f"J{row}")}'
            for row in range(99, 112)
            if text(f'B{row}')
        ]
        relationships = [
            f'{text(f"W{row}")}：{text(f"AD{row}")}'
            for row in range(131, 138)
            if text(f'W{row}')
        ]
        conditions = []
        for address in ('I11', 'R11'):
            state = text(address)
            if state and state not in ('健康', '清醒'):
                conditions.append(CONDITION_NAMES.get(state, state))
        inventory = []
        for row in range(79, 94):
            for address in (f'F{row}', f'N{row}'):
                name = text(address)
                if name and name not in ('0', ' '):
                    inventory.append({'name': name, 'source_cell': address})
        weapons = []
        for row in range(53, 58):
            name = text(f'B{row}')
            if name and name != '无':
                weapons.append(
                    {
                        'name': name,
                        'skill': text(f'M{row}'),
                        'damage': text(f'W{row}'),
                        'range': text(f'AA{row}'),
                        'attacks': text(f'AE{row}'),
                        'capacity': text(f'AG{row}'),
                        'malfunction': text(f'AJ{row}'),
                    }
                )
        max_hp = max(1, (attributes['CON'] + attributes['SIZ']) // 10)
        max_mp = attributes['POW'] // 5
        character = Character(
            actor_id=actor_id,
            name=text('E3', source_path.stem),
            occupation=text('E5'),
            age=number('E6', 30),
            sex=text('M6'),
            era=text('M4', '1920s'),
            residence=text('E7'),
            birthplace=text('M7'),
            attributes=attributes,
            skills=skills,
            skill_marks=marks,
            luck=number('AG7', 50),
            max_hp=max_hp,
            current_hp=number('E10', max_hp),
            max_mp=max_mp,
            current_mp=number('W10', max_mp),
            current_san=number('N10', attributes['POW']),
            background=background,
            experiences=experiences,
            relationships=relationships,
            inventory=inventory,
            weapons=weapons,
            conditions=conditions,
            allocations={'imported': imported_allocations},
            assets={
                'cash': text('O62'),
                'details': text('L63'),
                'currency': text('S62', '美元'),
            },
            source=f'CY23.5导入：{source_path.name}',
        )
        recalculated = derived_character(character)
        for name in ('current_hp', 'current_mp', 'current_san'):
            if getattr(character, name) != getattr(recalculated, name):
                warnings.append(f'{name}超过重算上限，已截断并需核对')
        namespace = '{http://schemas.openxmlformats.org/'
        namespace += 'spreadsheetml/2006/main}'
        with ZipFile(source_path) as archive:
            for part in archive.namelist():
                if not part.startswith('xl/worksheets/sheet') or not (
                    part.endswith('.xml')
                ):
                    continue
                tree = ElementTree.fromstring(archive.read(part))
                for item in tree.iter(namespace + 'c'):
                    formula = item.find(namespace + 'f')
                    if formula is not None and '#REF!' in (formula.text or ''):
                        warnings.append(
                            f'原表公式损坏：{part}!{item.attrib["r"]}'
                        )
        recalculated.import_warnings = list(dict.fromkeys(warnings))
        if any(
            background[name] in ('', '无') for name in BACKGROUND_FIELDS[:6]
        ):
            recalculated.import_warnings.append('核心背景缺失，请补充人物经历')
        return recalculated
    finally:
        formulas.close()
        values.close()
