from contextlib import contextmanager
from copy import deepcopy
from types import SimpleNamespace

import pytest
from openpyxl import Workbook

from aitrpg.adapters.cards import import_character_xlsx
from aitrpg.application.characters import CharacterAllocation
from aitrpg.application.characters import CharacterConcept
from aitrpg.application.characters import CharacterService
from aitrpg.application.platform import Platform
from aitrpg.config import Settings
from aitrpg.domain.models import Actor
from aitrpg.domain.models import Character
from aitrpg.domain.models import Provider
from aitrpg.domain.rules import BACKGROUND_FIELDS
from aitrpg.domain.rules import OCCUPATIONS
from aitrpg.domain.rules import validate_allocations


class MemoryStore:
    def __init__(self):
        self.data = {}

    def get(self, kind, entity_id):
        return deepcopy(self.data.get(kind, {}).get(entity_id))

    def list(self, kind):
        return list(deepcopy(self.data.get(kind, {})).values())

    def put(self, kind, entity_id, data):
        self.data.setdefault(kind, {})[entity_id] = deepcopy(data)

    @contextmanager
    def transaction(self):
        yield self


def make_character(actor_id='actor'):
    attributes = {
        'STR': 60,
        'CON': 60,
        'SIZ': 60,
        'DEX': 60,
        'APP': 60,
        'INT': 60,
        'POW': 60,
        'EDU': 60,
    }
    selected = [
        '会计',
        '估价',
        '汽车驾驶',
        '历史',
        '图书馆使用',
        '导航',
        '取悦',
        '说服',
    ]
    occupational = {
        '信用评级': 30,
        '会计': 30,
        '估价': 30,
        '汽车驾驶': 30,
        '历史': 30,
        '图书馆使用': 30,
        '导航': 30,
        '取悦': 20,
        '说服': 10,
    }
    interests = {'图书馆使用': 20, '侦查': 30, '聆听': 30, '潜行': 40}
    background = {key: '完整具体的人物背景' for key in BACKGROUND_FIELDS}
    background['key_connection'] = 'significant_people'
    return Character(
        actor_id=actor_id,
        name='林医生',
        occupation='古董商',
        occupation_definition=OCCUPATIONS['古董商'],
        allocations={
            'selected_skills': selected,
            'occupational': occupational,
            'interests': interests,
        },
        attributes=attributes,
        skills=validate_allocations(
            attributes,
            OCCUPATIONS['古董商'],
            selected,
            occupational,
            interests,
        ),
        background=background,
        max_hp=12,
        current_hp=7,
        max_mp=12,
        current_mp=4,
        current_san=39,
    )


def test_save_refuses_game_lock_and_stale_edits():
    store = MemoryStore()
    actor = Actor(id='actor', name='模型')
    store.put('actor', actor.id, actor.model_dump(mode='json'))
    service = CharacterService(store, None)
    first = service.save(make_character())
    first.name = '改名'
    second = service.save(first)
    assert second.version == 2
    with pytest.raises(ValueError, match='已被更新'):
        service.save(first)
    locked = second.model_dump(mode='json')
    locked['locked_game_id'] = 'ongoing-game'
    store.put('character', second.id, locked)
    with pytest.raises(ValueError, match='占用'):
        service.save(second)
    assert service.get(second.id).locked_game_id == 'ongoing-game'


def test_import_preserves_runtime_values_and_recalculates_formula_fields(
    tmp_path,
):
    path = tmp_path / '调查员.xlsx'
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '人物卡'
    for address in ('U3', 'U5', 'U7', 'AA3', 'AA5', 'AA7', 'AG3', 'AG5'):
        sheet[address] = 60
    sheet.merge_cells('E3:Q3')
    sheet['E3'] = '陈调查员'
    sheet['E6'] = 30
    sheet['E10'] = 4
    sheet['N10'] = 29
    sheet['W10'] = 2
    sheet['G10'] = '=999'
    sheet['P10'] = '=999'
    sheet['Y10'] = '=999'
    sheet['AB35'] = '侦查'
    sheet['AF35'] = 25
    sheet['AH35'] = 6
    sheet['AJ35'] = 30
    sheet['AL35'] = 10
    sheet['AN35'] = '=999'
    sheet['X35'] = '☑'
    workbook.save(path)
    imported = import_character_xlsx(path, 'actor')
    assert imported.name == '陈调查员'
    assert imported.current_hp == 4
    assert imported.current_mp == 2
    assert imported.current_san == 29
    assert imported.max_hp == 12
    assert imported.max_mp == 12
    assert imported.skills['侦查'] == 71
    assert imported.skill_marks == ['侦查']
    platform = Platform(Settings(data_dir=tmp_path / 'platform'))
    platform.save_actor(Actor(id='actor', name='被动模型', connection='mcp'))
    stored = platform.characters.import_xlsx(str(path), 'actor')
    assert platform.store.get('character', stored.id)['current_hp'] == 4
    assert platform.characters.list('actor')[0].current_san == 29


def test_blank_template_is_not_a_playable_character(tmp_path):
    workbook = Workbook()
    workbook.active.title = '人物卡'
    path = tmp_path / '空白.xlsx'
    workbook.save(path)
    with pytest.raises(ValueError, match='缺少数值'):
        import_character_xlsx(path, 'actor')


def test_profession_and_interest_pools_are_additive_and_separate():
    attributes = make_character().attributes
    definition = OCCUPATIONS['古董商']
    selected = [
        '会计',
        '估价',
        '汽车驾驶',
        '历史',
        '图书馆使用',
        '导航',
        '取悦',
        '说服',
    ]
    occupational = {
        '信用评级': 30,
        '会计': 30,
        '估价': 30,
        '汽车驾驶': 30,
        '历史': 30,
        '图书馆使用': 30,
        '导航': 30,
        '取悦': 20,
        '说服': 10,
    }
    skills = validate_allocations(
        attributes,
        definition,
        selected,
        occupational,
        {'图书馆使用': 20, '侦查': 30, '聆听': 30, '潜行': 40},
    )
    assert skills['图书馆使用'] == 70
    assert skills['侦查'] == 55
    assert skills['信用评级'] == 30
    with pytest.raises(ValueError, match='职业点数'):
        validate_allocations(
            attributes, definition, selected, {'会计': 1}, {'侦查': 120}
        )


async def test_generation_corrects_invalid_allocation_without_changing_dice(
    monkeypatch,
    tmp_path,
):
    actor = Actor(id='actor', name='模型', provider_id='provider')
    provider = Provider(
        id='provider', name='接口', base_url='https://example.test'
    )
    monkeypatch.setattr(
        'aitrpg.application.characters.generate_attributes',
        lambda: (dict(make_character().attributes), []),
    )
    background = dict(make_character().background)

    class ModelClient:
        calls = 0
        budgets = []

        async def generate(self, provider, system, context, output_type):
            self.calls += 1
            if output_type is CharacterConcept:
                value = CharacterConcept(
                    name='少年古董商',
                    age=19,
                    sex='女',
                    residence='上海',
                    birthplace='苏州',
                    occupation='古董商',
                    age_reductions={'SIZ': 5},
                    background=background,
                )
            else:
                self.budgets.append(context['occupational_budget'])
                values = {
                    '信用评级': 30,
                    '会计': 30,
                    '估价': 30,
                    '汽车驾驶': 30,
                    '历史': 20,
                    '图书馆使用': 20,
                    '导航': 20,
                    '取悦': 20,
                    '说服': 20,
                }
                if self.calls == 2:
                    values['会计'] = 1
                value = CharacterAllocation(
                    selected_skills=[
                        '会计',
                        '估价',
                        '汽车驾驶',
                        '历史',
                        '图书馆使用',
                        '导航',
                        '取悦',
                        '说服',
                    ],
                    occupational=values,
                    interests={
                        '侦查': 30,
                        '聆听': 30,
                        '潜行': 30,
                        '图书馆使用': 30,
                    },
                )
            return SimpleNamespace(value=value, usage=10)

    client = ModelClient()
    platform = Platform(Settings(data_dir=tmp_path / 'platform'), client)
    platform.save_provider(provider)
    platform.save_actor(actor)
    service = platform.characters
    character = await service.generate('actor', '上海少年古董商')
    assert client.calls == 3
    assert client.budgets == [220, 220]
    assert character.attributes['SIZ'] == 55
    assert character.attributes['EDU'] == 55
    assert character.max_hp == 11
    assert character.skills['图书馆使用'] == 70
    assert service.get(character.id).name == '少年古董商'
