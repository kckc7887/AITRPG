import io
import shutil
import zipfile
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook
from PIL import Image
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject
from pypdf.generic import DictionaryObject
from pypdf.generic import NameObject

from aitrpg.adapters.attachments import AttachmentError
from aitrpg.adapters.attachments import ExtractionBudget
from aitrpg.adapters.attachments import parse_attachment
from aitrpg.adapters.attachments import unpack_archive
from aitrpg.adapters.providers import Generation
from aitrpg.adapters.storage import Store
from aitrpg.application.scenarios import LinkCorrection
from aitrpg.application.scenarios import NavigationCorrections
from aitrpg.application.scenarios import ScenarioDraft
from aitrpg.application.scenarios import ScenarioService
from aitrpg.application.scenarios import _normalize_scene_participants
from aitrpg.application.scenarios import _normalize_source_endings
from aitrpg.application.scenarios import _normalize_terminal_links
from aitrpg.domain.models import Provider
from aitrpg.domain.models import Scenario
from aitrpg.domain.models import ScenarioContent
from aitrpg.domain.models import ScenarioRole
from aitrpg.domain.models import ScenarioScene
from aitrpg.domain.models import SourceBlock


def test_source_ending_list_keeps_transitions_out_of_terminal_catalog():
    first = SourceBlock(
        file='story.docx', locator='table:1', text='结束城内调查，进入第二章。'
    )
    ending_a = SourceBlock(
        file='story.docx', locator='table:8', text='结局A\n调查员平安离开。'
    )
    ending_b = SourceBlock(
        file='story.docx',
        locator='table:9',
        text='结局B\n调查员选择继续研究。',
    )
    scenario = Scenario(
        title='结局审查',
        source_blocks=[
            first,
            SourceBlock(
                file='story.docx', locator='paragraph:1', text='【结局列表】'
            ),
            ending_a,
            ending_b,
            SourceBlock(
                file='story.docx', locator='paragraph:2', text='【后记】'
            ),
        ],
        scenes=[ScenarioScene(id='city', title='城内', source_ids=[first.id])],
        endings=[
            ScenarioContent(
                id='city_done',
                title='城内调查结束',
                text='第二章还有后续。',
                source_ids=[first.id],
            ),
            ScenarioContent(
                id='ending_a', title='结局A：离开', source_ids=[ending_a.id]
            ),
            ScenarioContent(
                id='another_a', title='结局A：离开现场', source_ids=[first.id]
            ),
        ],
    )
    _normalize_source_endings(scenario)
    assert [ending.title for ending in scenario.endings] == ['结局A', '结局B']
    assert scenario.endings[0].text == '结局A\n调查员平安离开。'
    assert any('第二章还有后续' in item.text for item in scenario.handouts)
    assert all(item.visibility == 'keeper' for item in scenario.handouts)


def test_explicit_solo_and_split_party_source_scope_survives_import():
    first = SourceBlock(
        file='story.docx',
        locator='table:17:row:1:cell:1',
        text='这个部分只有ho1、ho3可参与。',
    )
    second = SourceBlock(
        file='story.docx',
        locator='table:17:row:1:cell:2:chars:6000-8000',
        text='深处的私密房间。',
    )
    scenario = Scenario(
        title='分队',
        source_blocks=[first, second],
        roles=[
            ScenarioRole(id='ho1', name='一'),
            ScenarioRole(id='ho2', name='二'),
            ScenarioRole(id='ho3', name='三'),
        ],
        scenes=[
            ScenarioScene(id='solo', title='HO2间章：私下谈话'),
            ScenarioScene(id='cave', title='洞穴深处', source_ids=[second.id]),
        ],
    )
    _normalize_scene_participants(scenario)
    assert scenario.scenes[0].conditions['participants'] == ['ho2']
    assert scenario.scenes[1].conditions['participants'] == ['ho1', 'ho3']


def test_metadata_can_arrive_after_preceding_portrait_chunks():
    scenario = Scenario(title='文件名')
    fields = set()
    ScenarioService._merge_draft(
        scenario, ScenarioDraft(author='作者'), 0, fields
    )
    ScenarioService._merge_draft(
        scenario,
        ScenarioDraft(
            title='正文标题',
            era='修仙纪元2026年',
            min_players=4,
            max_players=4,
        ),
        3,
        fields,
    )
    assert scenario.title == '正文标题'
    assert scenario.era == '修仙纪元2026年'
    assert scenario.min_players == scenario.max_players == 4


async def test_navigation_repair_accepts_only_existing_sourced_destinations(
    tmp_path: Path,
):
    class RepairProvider:
        def __init__(self, target):
            self.target = target

        async def generate(self, provider, system, context, output_type):
            return Generation(
                value=NavigationCorrections(
                    links=[
                        LinkCorrection(
                            scene_id='entrance',
                            old_target='legacy-end',
                            replacement_ids=[self.target],
                            source_ids=[context['sources'][0]['id']],
                            explanation='原文说明离开时进入结局A。',
                        )
                    ]
                ),
                usage=100,
            )

    store = Store(tmp_path / 'state.db')
    provider = Provider(name='test', base_url='http://localhost')
    store.put('provider', provider.id, provider.model_dump(mode='json'))
    source = SourceBlock(
        file='story.docx',
        locator='paragraph:1',
        text='调查员离开门厅时进入结局A。',
    )
    scenario = Scenario(
        title='导航',
        source_blocks=[source],
        scenes=[
            ScenarioScene(
                id='entrance',
                title='门厅',
                next_scene_ids=['legacy-end'],
                source_ids=[source.id],
            ),
        ],
        endings=[
            ScenarioContent(
                id='ending-a', title='结局A', source_ids=[source.id]
            )
        ],
    )
    malicious = ScenarioService(store, RepairProvider('invented'), tmp_path)
    await malicious._repair_navigation(scenario, provider.id)
    assert scenario.scenes[0].next_scene_ids == ['legacy-end']
    service = ScenarioService(store, RepairProvider('ending-a'), tmp_path)
    await service._repair_navigation(scenario, provider.id)
    assert scenario.scenes[0].next_scene_ids == []
    assert scenario.scenes[0].conditions['ending_ids'] == ['ending-a']
    assert scenario.scenes[0].conditions['import_navigation_rewrites'] == {
        'legacy-end': ['ending-a'],
    }
    store.close()


def test_terminal_links_and_repeated_ho_preserve_meaning():
    scenario = Scenario(
        title='分片',
        roles=[
            ScenarioRole(
                id='ho1',
                name='剑宗天才',
                secret_text='你由大师兄教导。',
            )
        ],
        scenes=[
            ScenarioScene(
                id='escape',
                title='逃离',
                next_scene_ids=['ending_safe'],
            )
        ],
        endings=[ScenarioContent(id='ending_safe', title='平安归来')],
    )
    later = ScenarioDraft(
        roles=[
            ScenarioRole(
                id='HO_1',
                name='命承灵眷',
                secret_text='你的灵根藏着真相。',
            )
        ]
    )
    ScenarioService._merge_draft(scenario, later, 1)
    assert len(scenario.roles) == 1
    assert '大师兄' in scenario.roles[0].secret_text
    assert '灵根' in scenario.roles[0].secret_text
    _normalize_terminal_links(scenario)
    assert scenario.scenes[0].next_scene_ids == []
    assert scenario.scenes[0].conditions['ending_ids'] == ['ending_safe']
    assert (
        'ending_safe'
        in (scenario.scenes[0].conditions['import_original_next_scene_ids'])
    )


def _pdf_text(path: Path, text: str) -> None:
    writer = PdfWriter()
    page = writer.add_blank_page(width=400, height=300)
    font = DictionaryObject(
        {
            NameObject('/Type'): NameObject('/Font'),
            NameObject('/Subtype'): NameObject('/Type1'),
            NameObject('/BaseFont'): NameObject('/Helvetica'),
        }
    )
    page[NameObject('/Resources')] = DictionaryObject(
        {
            NameObject('/Font'): DictionaryObject({NameObject('/F1'): font}),
        }
    )
    contents = DecodedStreamObject()
    contents.set_data(f'BT /F1 10 Tf 10 100 Td ({text}) Tj ET'.encode())
    page[NameObject('/Contents')] = writer._add_object(contents)
    writer.write(path)


def test_pdf_docx_duplicates_keep_one_body_but_versions_do_not(
    tmp_path: Path,
):
    source = tmp_path / 'input'
    source.mkdir()
    text = 'The lighthouse has a red lamp and a locked maintenance room.'
    document = Document()
    document.add_paragraph(text)
    document.save(source / 'story.docx')
    _pdf_text(source / 'story.pdf', text)
    result = parse_attachment(source, tmp_path / 'duplicate')
    assert sum(text in block.text for block in result.blocks) == 1
    assert any('重复' in issue.message for issue in result.issues)
    assert (tmp_path / 'duplicate' / 'raw' / 'story.pdf').is_file()
    document.add_paragraph(
        'The cellar contains an unknown machine. Its purpose is to send '
        'messages back in time. The new edition has a separate ending '
        'in which the heroes can choose to destroy or preserve the machine.'
    )
    document.save(source / 'story.docx')
    changed = parse_attachment(source, tmp_path / 'different')
    assert sum(text in block.text for block in changed.blocks) == 2
    assert any(issue.severity == 'blocker' for issue in changed.issues)


def test_docx_table_and_xlsx_formula_are_preserved_without_execution(
    tmp_path: Path,
):
    source = tmp_path / 'input'
    source.mkdir()
    document = Document()
    document.add_paragraph('第一章')
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = '玩家看到一盏灯。'
    table.cell(0, 1).text = 'KP 注释：灯中藏有怪物。'
    document.save(source / 'story.docx')
    workbook = Workbook()
    workbook.active['A1'] = 'STR'
    workbook.active['B1'] = '=SUM(40,20)'
    workbook.save(source / 'card.xlsx')
    result = parse_attachment(source, tmp_path / 'output')
    keeper = next(block for block in result.blocks if '藏有怪物' in block.text)
    assert keeper.locator == 'table:1:row:1:cell:2'
    assert any(block.text == '玩家看到一盏灯。' for block in result.blocks)
    formula = next(block for block in result.blocks if block.kind == 'formula')
    assert formula.text == '=SUM(40,20)'
    assert formula.locator.endswith(':B1')


@pytest.mark.parametrize('name', ['NUL.txt', '.. /outside.txt', 'notes.'])
def test_archive_rejects_windows_path_aliases(tmp_path: Path, name: str):
    archive_path = tmp_path / 'aliases.zip'
    with zipfile.ZipFile(archive_path, 'w') as archive:
        archive.writestr(name, '应保留的原文')
    with pytest.raises(AttachmentError, match='路径'):
        unpack_archive(archive_path, tmp_path / 'output')


def test_archive_rejects_traversal_links_and_size_limits(tmp_path: Path):
    archive_path = tmp_path / 'unsafe.zip'
    with zipfile.ZipFile(archive_path, 'w') as archive:
        archive.writestr('../outside.txt', 'secret')
    with pytest.raises(AttachmentError, match='路径'):
        unpack_archive(archive_path, tmp_path / 'output')
    assert not (tmp_path / 'outside.txt').exists()
    with zipfile.ZipFile(archive_path, 'w') as archive:
        link = zipfile.ZipInfo('link')
        link.create_system = 3
        link.external_attr = 0o120777 << 16
        archive.writestr(link, '../outside.txt')
    with pytest.raises(AttachmentError, match='链接'):
        unpack_archive(archive_path, tmp_path / 'output')
    with zipfile.ZipFile(archive_path, 'w') as archive:
        archive.writestr('a.txt', 'a')
    with pytest.raises(AttachmentError, match='总量'):
        unpack_archive(
            archive_path, tmp_path / 'output', ExtractionBudget(files=5000)
        )


def test_nested_zip_import_keeps_assets_and_does_not_run_code(tmp_path: Path):
    marker = tmp_path / 'executed.txt'
    image = io.BytesIO()
    Image.new('RGB', (12, 8), 'blue').save(image, format='PNG')
    nested = io.BytesIO()
    with zipfile.ZipFile(nested, 'w') as archive:
        archive.writestr('地图.png', image.getvalue())
        archive.writestr('readme.txt', '仅供同好交流，禁止商业用途。')
        archive.writestr('script.py', f'open({str(marker)!r}, "w").close()')
    path = tmp_path / 'module.zip'
    with zipfile.ZipFile(path, 'w') as archive:
        archive.writestr('nested.zip', nested.getvalue())
    result = parse_attachment(path, tmp_path / 'output')
    assert not marker.exists()
    assert any('禁止商业用途' in block.text for block in result.blocks)
    assert result.assets[0].is_map
    assert result.assets[0].visibility == 'keeper'
    assert (tmp_path / 'output' / result.assets[0].path).is_file()


class StructuredProvider:
    async def generate(self, provider, system, context, output_type):
        source = context['sources'][0]['id']
        value = ScenarioDraft(
            title='私人剧本',
            scenes=[
                ScenarioScene(
                    id='arrival',
                    title='门厅',
                    public_text='灯亮了。',
                    keeper_text='一切是幕后机器所为。',
                    source_ids=[source],
                )
            ],
            roles=[
                ScenarioRole(
                    id='ho1',
                    name='医生',
                    secret_text='你曾在此工作。',
                    source_ids=[source],
                )
            ],
        )
        return Generation(value=value, usage=100)


async def test_conversion_and_approved_version_survive_later_edits(
    tmp_path: Path,
):
    source = tmp_path / 'story.docx'
    document = Document()
    document.add_paragraph('门厅：灯亮了；幕后机器；医生曾在此工作。')
    document.save(source)
    store = Store(tmp_path / 'state.db')
    provider = Provider(name='test', base_url='http://localhost')
    store.put('provider', provider.id, provider.model_dump(mode='json'))
    service = ScenarioService(store, StructuredProvider(), tmp_path / 'data')
    draft = await service.import_path(source, provider.id)
    approved = service.approve(draft.id)
    approved.scenes[0].keeper_text = '新的真相'
    edited = service.save(approved)
    assert edited.status == 'draft'
    assert edited.version == 2
    frozen = store.get('scenario_version', f'{approved.id}@1')
    assert frozen['scenes'][0]['keeper_text'] == '一切是幕后机器所为。'
    assert service.get(approved.id).scenes[0].keeper_text == '新的真相'
    store.close()


async def test_canonical_package_round_trip_and_immutable_map(tmp_path: Path):
    store = Store(tmp_path / 'state.db')
    service = ScenarioService(store, None, tmp_path / 'data')
    original = service.seed_demo()
    package = service.export_package(original.id)
    imported = await service.import_path(package)
    assert imported.id != original.id
    assert imported.status == 'draft'
    assert imported.scenes[0].keeper_text == original.scenes[0].keeper_text
    assert imported.assets[0].regions == original.assets[0].regions
    frozen = Scenario.model_validate(
        store.get(
            'scenario_version',
            f'{original.id}@1',
        )
    )
    original_path = Path(original.directory) / original.assets[0].path
    old_path = Path(frozen.directory) / frozen.assets[0].path
    old_bytes = old_path.read_bytes()
    Image.new('RGB', (600, 300), 'yellow').save(original_path)
    assert old_path.read_bytes() == old_bytes
    assert old_path != original_path
    with zipfile.ZipFile(package) as archive:
        assert not any(name.startswith('raw/') for name in archive.namelist())
    store.close()


def test_missing_rar_backend_is_explicit(tmp_path: Path, monkeypatch):
    from aitrpg.adapters import attachments

    monkeypatch.delenv('AITRPG_BSDTAR', raising=False)
    monkeypatch.setattr(shutil, 'which', lambda name: None)
    path = tmp_path / 'module.rar'
    path.write_bytes(b'not-real-rar')
    with pytest.raises(AttachmentError, match='bsdtar'):
        attachments.unpack_archive(path, tmp_path / 'output')
